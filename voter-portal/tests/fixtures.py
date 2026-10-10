"""Build four small, entirely made-up source files that exercise the messy cases."""
import csv
import io
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from PIL import Image


def _png(color):
    buf = io.BytesIO()
    Image.new("RGB", (40, 50), color).save(buf, "PNG")
    buf.seek(0)
    return buf


def build(data_dir):
    data_dir = Path(data_dir)
    data_dir.mkdir(parents=True, exist_ok=True)
    (data_dir / "pics").mkdir(exist_ok=True)

    # 1. Gahdhdhoo.2026 — title rows above the header, embedded pictures, a repeated header.
    wb = Workbook()
    ws = wb.active
    ws.title = "List"
    ws.append(["GADHDHOO VOTERS 2026"])
    ws.append([])
    ws.append(["#", "Photo", "ID Card No.", "Name", "Gender", "Age", "Address", "Party"])
    ws.append([1, None, "A100001", "Aishath Test", "Female", 34, "Test House", "AAA"])   # row 4
    ws.append([2, None, "a 100002", "Ibrahim Sample", "M", "52", "Sample Villa", "BBB"])  # row 5
    ws.append(["#", "Photo", "ID Card No.", "Name", "Gender", "Age", "Address", "Party"])  # repeated header
    ws.append([3, None, "A100003", "Mariyam Fake", "އަންހެން", 17, "Fake Ge", ""])        # row 7
    ws.add_image(XLImage(_png("red")), "B4")
    ws.add_image(XLImage(_png("blue")), "B5")
    wb.save(data_dir / "Gahdhdhoo.2026.xlsx")

    # 2. gadhoo-new voters — CSV, numeric ID missing the A, photo file-name column.
    Image.new("RGB", (30, 30), "green").save(data_dir / "pics" / "IMG_0004.jpg")
    with open(data_dir / "gadhoo-new voters.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["National ID", "Full Name", "Sex", "Age", "House Name", "Party", "Photo File"])
        w.writerow(["100004", "Hassan  Demo ", "m", "19", "Demo Hiya", "ccc", r"C:\Users\x\pics\IMG_0004.jpg"])
        w.writerow(["A100001", "Aishath Test", "F", "33", "Test House", "", ""])
        w.writerow(["", "No Id Person", "F", "40", "Nowhere", "AAA", ""])

    # 3. Updated party list — two party columns, conflicting details.
    wb = Workbook()
    ws = wb.active
    ws.append(["ID", "Name", "Sex", "Age", "House", "Party", "Updated Party"])
    ws.append(["A100002", "Ibrahim Sample", "M", 52, "Sample Villa", "BBB", "CCC"])
    ws.append(["A100005", "Ali Placeholder", "M", 45, "Old House", "BBB", "AAA"])
    ws.append(["A100006", "Totally Different", "F", 60, "Other", "BBB", "BBB"])
    wb.save(data_dir / "Gadhdhoo (updated Party list MPD ).xlsx")

    # 4. Full community list — oldest data, duplicates and a likely duplicate person.
    wb = Workbook()
    ws = wb.active
    ws.append(["id card", "name", "gender", "age", "house name", "party"])
    ws.append(["A100005", "Ali Placeholder", "M", 44, "New House", "BBB"])
    ws.append(["A100006", "Zaid Unrelated", "M", 30, "Other", "BBB"])
    ws.append(["A100007", "Hawwa Dup", "F", 70, "Same Ge", "AAA"])
    ws.append(["A100008", "Hawwa  Dup", "F", 70, "Same Ge", "AAA"])
    ws.append(["A100007", "Hawwa Dup", "F", 70, "Same Ge", "AAA"])
    ws.append(["A100009", "Photo ByName", "F", 28, "Pic House", "AAA"])
    wb.save(data_dir / "final_cleaned_gadhdhoo_updated_party.xlsx")
    Image.new("RGB", (30, 30), "yellow").save(data_dir / "pics" / "A100009.png")
    return data_dir
