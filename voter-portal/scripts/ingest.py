#!/usr/bin/env python3
"""Merge the Gadhdhoo voter files into one de-duplicated list and load it.

    python scripts/ingest.py --data-dir "/path/to/Data files"            # merge + load
    python scripts/ingest.py --data-dir "/path/to/Data files" --dry-run  # report only

What it does
  1. Finds the four source files in --data-dir (any of .xlsx .xlsm .xls .csv; the
     names are matched loosely, so "Gadhdhoo (updated Party list MPD ).xlsx" is found
     even with odd spacing or case).
  2. Finds the header row and the ID / name / sex / age / house / party / photo
     columns in every sheet, by header name first and by the values second (a column
     where most cells look like A123456 is the ID column, and so on).
  3. Cleans every value (whitespace, ID format, sex, age, party upper-case).
  4. De-duplicates strictly on the national ID. For each field the value from the
     highest-priority file that has one wins:
        Gahdhdhoo.2026  >  gadhoo-new voters  >  Gadhdhoo (updated Party list MPD )
                        >  final_cleaned_gadhdhoo_updated_party
     and every disagreement is written to merge_conflicts.log.
  5. Takes photos from pictures embedded in the Gahdhdhoo.2026 workbook (matched to
     the row they sit on), from a photo/file-name column, or from image files named
     after the ID card (A123456.jpg) anywhere under --data-dir.
  6. Creates the tables, loads the voters and seeds an admin and a standard user.

merge_conflicts.log and --csv-out contain personal data: keep them on the server.
"""
import argparse
import csv
import difflib
import io
import os
import re
import secrets
import sys
import threading
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path, PureWindowsPath

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.db import connect, init_schema, now_iso  # noqa: E402
from app.normalize import (  # noqa: E402
    ID_RE,
    clean_text,
    compare_key,
    is_blank,
    normalize_age,
    normalize_id,
    normalize_party,
    normalize_sex,
)

# Highest priority first. `key` is matched against file names loosely.
SOURCES = [
    ("Gahdhdhoo.2026", "2026 list (with photos)"),
    ("gadhoo-new voters", "new voters"),
    ("Gadhdhoo (updated Party list MPD )", "updated party list"),
    ("final_cleaned_gadhdhoo_updated_party", "full community list"),
]
TABLE_EXTS = {".xlsx", ".xlsm", ".xls", ".csv"}
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp"}
FIELDS = ("full_name", "sex", "age", "house_name", "party")

# Header names, compared after lower-casing and removing everything but letters/digits.
ALIASES = {
    "national_id": {"id", "idcard", "idcardno", "idcardnumber", "idno", "idnumber", "nid", "nationalid",
                    "nationalidcard", "nationalidno", "nationalidnumber", "identitycard", "identitycardno",
                    "nic", "idcardnum", "voterid"},
    "full_name": {"name", "fullname", "votername", "nameofvoter", "names", "englishname", "fullnameenglish",
                  "nameenglish", "voter"},
    "sex": {"sex", "gender", "mf"},
    "age": {"age", "ages", "ageyears", "ageyrs"},
    "dob": {"dob", "dateofbirth", "birthdate", "birthday"},
    "house_name": {"house", "housename", "address", "houseaddress", "permanentaddress", "home", "residence",
                   "housenameaddress", "homename"},
    "party": {"party", "partyname", "politicalparty", "affiliation", "partyaffiliation"},
    "photo": {"photo", "photopath", "photofile", "photoname", "image", "imagefile", "picture", "pic",
              "photoref", "img", "filename", "file"},
}
NEWER_WORDS = ("updat", "new", "latest", "current", "2026")
# Fallback substring rules, checked in this order.
SUBSTRINGS = [
    ("photo", ("photo", "image", "picture")),
    ("house_name", ("house", "address")),
    ("party", ("party",)),
    ("national_id", ("idcard", "nationalid", "identity")),
    ("sex", ("gender",)),
    ("dob", ("birth",)),
    ("full_name", ("name",)),
]


_local = threading.local()


class IngestError(Exception):
    """A problem with the input that stops the import (shown to the user as is)."""


def say(*args):
    print(*args, file=getattr(_local, "out", None) or sys.stdout)


def header_key(text):
    return re.sub(r"[\W_]+", "", clean_text(text).casefold())


def role_for_header(text):
    key = header_key(text)
    if not key:
        return None
    for role, names in ALIASES.items():
        if key in names:
            return role
    for role, needles in SUBSTRINGS:
        if any(n in key for n in needles):
            return role
    return None


# --- conflict log -------------------------------------------------------------

class ConflictLog:
    def __init__(self):
        self.lines = []
        self.counts = Counter()

    def add(self, level, code, nid, message):
        self.counts[(level, code)] += 1
        self.lines.append(f"{level:<5} {code:<26} {nid or '-':<10} {message}")

    def write(self, path, header):
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        with open(path, "w", encoding="utf-8") as fh:
            fh.write("\n".join(header) + "\n\n")
            fh.write("Summary\n")
            for (level, code), n in sorted(self.counts.items()):
                fh.write(f"  {level:<5} {code:<26} {n}\n")
            fh.write("\nDetails (LEVEL CODE ID MESSAGE)\n")
            fh.write("\n".join(self.lines) + "\n")
        try:
            os.chmod(path, 0o600)
        except OSError:
            pass


# --- reading files --------------------------------------------------------------

@dataclass
class Sheet:
    name: str
    rows: list
    images: list = field(default_factory=list)  # (first row, last row, bytes) per embedded picture


def read_xlsx(path):
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)  # not read-only, so embedded images load
    sheets = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        images = []
        for img in getattr(ws, "_images", []):
            start = getattr(getattr(img.anchor, "_from", None), "row", None)
            if start is None:
                continue
            end = getattr(getattr(img.anchor, "to", None), "row", start)
            try:
                images.append((start, max(start, end), img._data()))
            except Exception:
                continue
        sheets.append(Sheet(ws.title, rows, images))
    return sheets


def read_xls(path):
    import pandas as pd

    frames = pd.read_excel(path, sheet_name=None, header=None, dtype=object)
    return [Sheet(name, df.where(df.notna(), None).values.tolist()) for name, df in frames.items()]


def read_csv(path):
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            with open(path, newline="", encoding=encoding) as fh:
                sample = fh.read(4096)
                fh.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                except csv.Error:
                    dialect = csv.excel
                return [Sheet("csv", [row for row in csv.reader(fh, dialect)])]
        except UnicodeError:
            continue
    raise ValueError(f"Could not decode {path}")


def read_table(path):
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        return read_xlsx(path)
    if ext == ".xls":
        return read_xls(path)
    return read_csv(path)


# --- column detection -------------------------------------------------------------

def find_header(rows, scan=30):
    """Return (header row index, {role: column}) or (-1, {}) if no header found."""
    best = (-1, {})
    for i, row in enumerate(rows[:scan]):
        candidates = defaultdict(list)
        for c, cell in enumerate(row):
            role = role_for_header(cell) if isinstance(cell, str) else None
            if role:
                candidates[role].append(c)
        roles = {}
        for role, cols in candidates.items():
            # Two "Party" columns (old / updated): take the newest-sounding one.
            fresh = [c for c in cols if any(w in header_key(row[c]) for w in NEWER_WORDS)]
            roles[role] = (fresh or cols)[0]
        if len(roles) >= 2 and len(roles) > len(best[1]):
            best = (i, roles)
    return best


def sniff_columns(rows, start, roles):
    """Fill in ID / sex / age columns from the values when the header didn't name them."""
    body = rows[start:start + 400]
    width = max((len(r) for r in body), default=0)
    taken = set(roles.values())

    def share(col, test):
        vals = [r[col] for r in body if col < len(r) and not is_blank(r[col])]
        if len(vals) < 3:
            return 0
        if test is looks_like_age:
            nums = [float(str(v).strip()) for v in vals if looks_like_age(v)]
            if nums == sorted(nums) and len(set(nums)) == len(nums):
                return 0  # 1, 2, 3, ... is a serial-number column, not ages
        return sum(1 for v in vals if test(v)) / len(vals)

    def looks_like_id(v):
        return bool(ID_RE.match(re.sub(r"[\s\-]", "", str(v).upper())))

    def looks_like_sex(v):
        return normalize_sex(v)[0] != ""

    def looks_like_age(v):
        try:
            return 0 < float(str(v).strip()) <= 120 and float(str(v).strip()).is_integer()
        except ValueError:
            return False

    for role, test, threshold in (("national_id", looks_like_id, 0.6),
                                  ("sex", looks_like_sex, 0.8),
                                  ("age", looks_like_age, 0.8)):
        if role in roles or (role == "age" and "dob" in roles):
            continue
        scores = [(share(c, test), c) for c in range(width) if c not in taken]
        if scores:
            score, col = max(scores)
            if score >= threshold:
                roles[role] = col
                taken.add(col)
    return roles


def age_from_dob(value):
    if isinstance(value, datetime):
        born = value.date()
    elif isinstance(value, date):
        born = value
    else:
        text = clean_text(value)
        born = None
        for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%Y", "%d %b %Y", "%d %B %Y"):
            try:
                born = datetime.strptime(text, fmt).date()
                break
            except ValueError:
                continue
        if born is None:
            return None
    today = date.today()
    return today.year - born.year - ((today.month, today.day) < (born.month, born.day))


# --- photos -----------------------------------------------------------------------

class PhotoFinder:
    """Index of image files under the data folder, by lower-case file name and by stem."""

    def __init__(self, data_dir):
        self.by_name, self.by_stem = {}, {}
        for p in Path(data_dir).rglob("*"):
            if p.is_file() and p.suffix.lower() in IMAGE_EXTS:
                self.by_name.setdefault(p.name.casefold(), p)
                self.by_stem.setdefault(p.stem.casefold(), p)

    def lookup(self, reference):
        text = clean_text(reference)
        if not text:
            return None
        name = PureWindowsPath(text).name if "\\" in text else Path(text).name
        direct = Path(text)
        if direct.is_file():
            return direct
        return self.by_name.get(name.casefold()) or self.by_stem.get(Path(name).stem.casefold())

    def for_id(self, nid):
        return self.by_stem.get(nid.casefold())


def store_photo(blob, nid, photos_dir):
    """Re-encode as a JPEG (max 800px, no EXIF) and return the stored file name."""
    from PIL import Image

    img = Image.open(io.BytesIO(blob))
    img.load()
    img = img.convert("RGB")
    img.thumbnail((800, 800))
    name = f"{nid}.jpg"
    img.save(Path(photos_dir) / name, "JPEG", quality=85)
    return name


# --- extraction -------------------------------------------------------------------

@dataclass
class Record:
    nid: str
    priority: int
    source: str
    where: str
    values: dict
    photo_blob: bytes = None
    photo_ref: str = None


def extract(path, priority, label, log, finder):
    records, stats = [], Counter()
    for sheet in read_table(path):
        header_idx, roles = find_header(sheet.rows)
        roles = sniff_columns(sheet.rows, header_idx + 1, roles)
        if "national_id" not in roles:
            if any(any(not is_blank(c) for c in r) for r in sheet.rows):
                log.add("WARN", "SHEET_SKIPPED", None, f"{label} / sheet '{sheet.name}': no ID card column found")
            continue
        missing = [r for r in ("full_name", "sex", "house_name", "party") if r not in roles]
        stats["sheets"] += 1
        say(f"  sheet '{sheet.name}': header row {header_idx + 1 if header_idx >= 0 else 'none'}, "
              f"columns {sorted(roles)}" + (f", missing {missing}" if missing else ""))
        header_keys = [header_key(c) for c in sheet.rows[header_idx]] if header_idx >= 0 else None

        def cell(row, role):
            col = roles.get(role)
            return row[col] if col is not None and col < len(row) else None

        row_records = {}
        for i in range(header_idx + 1, len(sheet.rows)):
            row = sheet.rows[i]
            if all(is_blank(c) for c in row):
                continue
            if header_keys and [header_key(c) for c in row] == header_keys:
                continue  # header repeated on every page of a PDF export
            where = f"{label} / {sheet.name} row {i + 1}"
            nid, note = normalize_id(cell(row, "national_id"))
            if nid is None:
                stats["skipped"] += 1
                name = clean_text(cell(row, "full_name"))
                log.add("WARN", "NO_VALID_ID", None, f"{where}: {note}; row skipped" + (f" (name '{name}')" if name else ""))
                continue
            if note:
                log.add("INFO", "ID_NORMALISED", nid, f"{where}: {note}")
            values = {"full_name": clean_text(cell(row, "full_name")), "house_name": clean_text(cell(row, "house_name")),
                      "party": normalize_party(cell(row, "party"))}
            values["sex"], note = normalize_sex(cell(row, "sex"))
            if note:
                log.add("WARN", "BAD_SEX", nid, f"{where}: {note}")
            if "age" in roles:
                values["age"], note = normalize_age(cell(row, "age"))
                if note:
                    log.add("WARN", "BAD_AGE", nid, f"{where}: {note}")
            elif "dob" in roles:
                values["age"] = age_from_dob(cell(row, "dob")) if not is_blank(cell(row, "dob")) else None
            else:
                values["age"] = None
            rec = Record(nid, priority, label, where, values)
            if "photo" in roles and not is_blank(cell(row, "photo")):
                ref = clean_text(cell(row, "photo"))
                found = finder.lookup(ref)
                if found:
                    rec.photo_ref = str(found)
                else:
                    log.add("WARN", "PHOTO_NOT_FOUND", nid, f"{where}: photo '{ref}' not found under the data folder")
            records.append(rec)
            row_records[i] = rec
            stats["rows"] += 1
        attach_pictures(sheet, row_records, label, log)
    return records, stats


def attach_pictures(sheet, row_records, label, log):
    """Give each embedded picture to the voter row it sits on (top row first, then any row it covers)."""
    orphans = []
    for start, end, blob in sorted(sheet.images, key=lambda x: x[0]):
        row = start if start in row_records else next((r for r in range(start, end + 1) if r in row_records), None)
        if row is None:
            orphans.append(start)
            continue
        rec = row_records[row]
        if rec.photo_blob is not None:
            log.add("WARN", "MULTIPLE_PHOTOS", rec.nid, f"{rec.where}: more than one picture on this row; kept the first")
            continue
        rec.photo_blob = blob
        rec.photo_ref = None  # an embedded picture beats a file-name reference
    if orphans:
        log.add("WARN", "UNMATCHED_PICTURES", None,
                f"{label} / sheet '{sheet.name}': {len(orphans)} embedded pictures not on a voter row "
                f"(rows {', '.join(str(r + 1) for r in orphans[:20])})")


# --- merging ----------------------------------------------------------------------

def merge(records, log):
    by_id = defaultdict(list)
    for r in records:
        by_id[r.nid].append(r)

    merged = {}
    for nid, recs in by_id.items():
        recs.sort(key=lambda r: r.priority)  # stable: keeps file order within a source
        per_source = Counter(r.source for r in recs)
        for source, n in per_source.items():
            if n > 1:
                same = [r for r in recs if r.source == source]
                differs = any(a.values != same[0].values for a in same[1:])
                log.add("WARN" if differs else "INFO", "DUPLICATE_IN_FILE", nid,
                        f"{n} rows in {source} ({'different details' if differs else 'identical'}): "
                        + "; ".join(r.where.split(" / ", 1)[1] for r in same))

        out = {"national_id": nid}
        for f in FIELDS:
            chosen = next((r for r in recs if r.values.get(f) not in ("", None)), None)
            out[f] = chosen.values[f] if chosen else ("" if f != "age" else None)
            if not chosen:
                continue
            for other in recs:
                val = other.values.get(f)
                if other is chosen or val in ("", None):
                    continue
                check_conflict(f, nid, chosen, other, log)
        out["sources"] = sorted({r.source for r in recs}, key=lambda s: [x[1] for x in SOURCES].index(s))
        out["_photo_rec"] = next((r for r in recs if r.photo_blob or r.photo_ref), None)
        merged[nid] = out
    return merged


def check_conflict(f, nid, chosen, other, log):
    a, b = chosen.values[f], other.values[f]
    keep = f"kept '{a}' ({chosen.source}) over '{b}' ({other.source})"
    if f == "full_name":
        ka, kb = compare_key(a), compare_key(b)
        if ka == kb:
            return
        ratio = difflib.SequenceMatcher(None, ka, kb).ratio()
        if ratio < 0.6:
            log.add("WARN", "AMBIGUOUS_NAME_MISMATCH", nid,
                    f"names very different ({ratio:.0%} similar), the ID may be mistyped in one file; {keep}")
        else:
            log.add("INFO", "NAME_VARIANT", nid, keep)
    elif f == "age":
        if abs(a - b) > 1:  # one year apart is just an older list
            log.add("WARN", "AGE_CONFLICT", nid, keep)
    elif f == "sex":
        if a != b:
            log.add("WARN", "SEX_CONFLICT", nid, keep)
    elif f == "house_name":
        if compare_key(a) != compare_key(b):
            log.add("INFO", "HOUSE_CHANGED", nid, keep)
    elif f == "party":
        if a != b:
            log.add("INFO", "PARTY_CHANGED", nid, keep)


def post_checks(merged, log):
    groups = defaultdict(list)
    for nid, v in merged.items():
        if not v["full_name"]:
            log.add("WARN", "MISSING_NAME", nid, f"no name in any file ({', '.join(v['sources'])})")
        else:
            groups[(compare_key(v["full_name"]), compare_key(v["house_name"]))].append(nid)
        if v["age"] is not None and v["age"] < 18:
            log.add("WARN", "UNDER_18", nid, f"age {v['age']}")
    for (_name, _house), ids in groups.items():
        if len(ids) > 1:
            v = merged[ids[0]]
            log.add("WARN", "POSSIBLE_DUPLICATE_PERSON", ",".join(sorted(ids)),
                    f"same name and house '{v['full_name']}' / '{v['house_name']}' under {len(ids)} IDs")


# --- loading ----------------------------------------------------------------------

def resolve_photo(v, finder, photos_dir, log, dry_run):
    rec = v.pop("_photo_rec")
    blob = None
    if rec and rec.photo_blob:
        blob = rec.photo_blob
    elif rec and rec.photo_ref:
        blob = Path(rec.photo_ref).read_bytes()
    else:
        p = finder.for_id(v["national_id"])
        if p:
            blob = p.read_bytes()
    if blob is None:
        return None
    if dry_run:
        return f"{v['national_id']}.jpg"
    try:
        return store_photo(blob, v["national_id"], photos_dir)
    except Exception as exc:
        log.add("WARN", "BAD_PHOTO", v["national_id"], f"could not read picture: {exc}")
        return None


def load(conn, merged, mode, log):
    stamp = now_iso()
    counts = Counter()
    if mode == "replace":
        counts["deleted"] = conn.execute("DELETE FROM voters").rowcount
    for nid, v in merged.items():
        existing = conn.execute("SELECT * FROM voters WHERE national_id = ?", (nid,)).fetchone()
        src = "import: " + ", ".join(v["sources"])
        if existing is None:
            conn.execute(
                "INSERT INTO voters (national_id, full_name, sex, age, house_name, party, photo_path, source,"
                " created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (nid, v["full_name"], v["sex"], v["age"], v["house_name"], v["party"], v["photo_path"], src,
                 stamp, stamp),
            )
            counts["inserted"] += 1
        elif existing["updated_by"] is not None or existing["source"] == "manual":
            log.add("INFO", "KEPT_PORTAL_EDIT", nid, "edited in the portal since the last import; file data not applied")
            counts["kept_portal_edits"] += 1
        else:
            conn.execute(
                "UPDATE voters SET full_name = ?, sex = ?, age = ?, house_name = ?, party = ?,"
                " photo_path = COALESCE(?, photo_path), source = ?, updated_at = ? WHERE national_id = ?",
                (v["full_name"], v["sex"], v["age"], v["house_name"], v["party"], v["photo_path"], src, stamp, nid),
            )
            counts["updated"] += 1
    return counts


def seed_users(conn, admin_name, user_name):
    from app.auth import hash_password

    created = []
    for username, role, env in ((admin_name, "admin", "SEED_ADMIN_PASSWORD"), (user_name, "user", "SEED_USER_PASSWORD")):
        if conn.execute("SELECT 1 FROM users WHERE username = ?", (username,)).fetchone():
            say(f"  user '{username}' already exists; left unchanged")
            continue
        password = os.environ.get(env) or secrets.token_urlsafe(12)
        stamp = now_iso(precise=True)
        conn.execute(
            "INSERT INTO users (username, password_hash, role, must_change_password, password_changed_at, created_at)"
            " VALUES (?, ?, ?, 1, ?, ?)",
            (username, hash_password(password), role, stamp, stamp),
        )
        created.append((username, role, password, bool(os.environ.get(env))))
    return created


def find_sources(data_dir, log):
    candidates = [p for p in Path(data_dir).iterdir() if p.is_file() and p.suffix.lower() in TABLE_EXTS
                  and not p.name.startswith("~$")]
    by_key = {header_key(p.stem): p for p in candidates}
    found, used = [], set()
    for priority, (key, label) in enumerate(SOURCES):
        k = header_key(key)
        path = by_key.get(k)
        if path is None:
            close = difflib.get_close_matches(k, [x for x in by_key if by_key[x] not in used], n=1, cutoff=0.85)
            path = by_key[close[0]] if close else None
        if path is None or path in used:
            log.add("WARN", "SOURCE_MISSING", None, f"no file found for '{key}' ({label})")
            say(f"[!] missing: {key}")
            continue
        used.add(path)
        found.append((priority, label, path))
    for p in candidates:
        if p not in used:
            say(f"[i] ignored (not one of the four sources): {p.name}")
    return found


def default_paths():
    base = Path(os.environ.get("VOTER_PORTAL_DATA_DIR", ROOT / "instance"))
    return (Path(os.environ.get("VOTER_PORTAL_DB", base / "voter_portal.db")),
            Path(os.environ.get("VOTER_PORTAL_PHOTOS", base / "photos")))


def main(argv=None, out=None):
    _local.out = out
    db_default, photos_default = default_paths()
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--data-dir", required=True, help="folder holding the four voter files (and any photos)")
    ap.add_argument("--db", default=str(db_default), help=f"SQLite file (default {db_default})")
    ap.add_argument("--photos-dir", default=str(photos_default), help=f"where photos are stored (default {photos_default})")
    ap.add_argument("--log", default="merge_conflicts.log", help="conflict / anomaly log (default ./merge_conflicts.log)")
    ap.add_argument("--csv-out", help="also write the merged list to this CSV for review")
    ap.add_argument("--mode", choices=("upsert", "replace"), default="upsert",
                    help="upsert (default): add new IDs, refresh imported ones, keep voters edited in the portal; "
                         "replace: wipe the voters table first")
    ap.add_argument("--dry-run", action="store_true", help="merge and write the log, but don't touch the database")
    ap.add_argument("--admin-username", default="admin")
    ap.add_argument("--user-username", default="user")
    ap.add_argument("--no-seed-users", action="store_true")
    args = ap.parse_args(argv)

    data_dir = Path(args.data_dir)
    if not data_dir.is_dir():
        raise IngestError(f"{data_dir} is not a folder")
    log = ConflictLog()
    sources = find_sources(data_dir, log)
    if not sources:
        raise IngestError("None of the four voter files were found (supported: .xlsx .xlsm .xls .csv).")

    finder = PhotoFinder(data_dir)
    records, source_lines = [], []
    for priority, label, path in sources:
        say(f"[{priority + 1}] {label}: {path.name}")
        recs, stats = extract(path, priority, label, log, finder)
        records += recs
        source_lines.append(f"  priority {priority + 1}: {path.name} ({label}) - {stats['rows']} rows read, "
                            f"{stats['skipped']} skipped, {stats['sheets']} sheet(s)")

    merged = merge(records, log)
    post_checks(merged, log)

    photos_dir = Path(args.photos_dir)
    if not args.dry_run:
        photos_dir.mkdir(parents=True, exist_ok=True)
    for v in merged.values():
        v["photo_path"] = resolve_photo(v, finder, photos_dir, log, args.dry_run)

    with_photo = sum(1 for v in merged.values() if v["photo_path"])
    summary = [
        f"Gadhdhoo voter merge - {datetime.now():%Y-%m-%d %H:%M}",
        f"Data folder: {data_dir.name}",
        "Sources (highest priority first):",
        *source_lines,
        f"Rows read: {len(records)}  ->  unique voters: {len(merged)}  "
        f"(duplicates merged: {len(records) - len(merged)}); with photo: {with_photo}",
    ]

    if args.csv_out:
        with open(args.csv_out, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh)
            w.writerow(["ID", "full_name", "sex", "age", "house_name", "Party", "photo_path", "sources"])
            for nid in sorted(merged):
                v = merged[nid]
                w.writerow([nid, v["full_name"], v["sex"], v["age"], v["house_name"], v["party"],
                            v["photo_path"] or "", "; ".join(v["sources"])])
        os.chmod(args.csv_out, 0o600)

    if args.dry_run:
        summary.append("Dry run: database not changed.")
    else:
        Path(args.db).parent.mkdir(parents=True, exist_ok=True)
        conn = connect(args.db)
        init_schema(conn)
        with conn:
            counts = load(conn, merged, args.mode, log)
            created = [] if args.no_seed_users else seed_users(conn, args.admin_username, args.user_username)
            conn.execute(
                "INSERT INTO audit_log (at, username, action, target, detail) VALUES (?, 'ingest-script', 'import', ?, ?)",
                (now_iso(), args.mode, f"{len(merged)} voters; {dict(counts)}"),
            )
        conn.close()
        summary.append(f"Database {args.db} ({args.mode}): {dict(counts)}")

    log.write(args.log, summary)
    say("\n".join(summary))
    warns = sum(n for (lvl, _), n in log.counts.items() if lvl == "WARN")
    say(f"Conflict log: {args.log} ({warns} warnings, {len(log.lines) - warns} notes)")
    if not args.dry_run and not args.no_seed_users:
        for username, role, password, from_env in created:
            shown = "(from environment)" if from_env else password
            say(f"Seeded {role} '{username}' password: {shown}  - must be changed at first sign-in")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except IngestError as exc:
        sys.exit(f"error: {exc}")
